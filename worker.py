import asyncio
import base64
import json
import mimetypes
import os
import re
import sys
import time
from html import escape, unescape
from pathlib import Path
from urllib.parse import quote_plus, urljoin, urlparse
import xml.etree.ElementTree as ET

from nats.aio.client import Client as NATS
from openai import AsyncOpenAI
import redis.asyncio as redis


# ==========================================
# CONFIGURATION
# ==========================================
NATS_URL = os.getenv("NATS_URL", "nats://localhost:4222")
REDIS_URL = os.getenv("REDIS_URL", "redis://localhost:6379")
LOCAL_LLM_URL = os.getenv("LOCAL_LLM_URL", "http://localhost:11434/v1")
LOCAL_LLM_MODEL = os.getenv("LOCAL_LLM_MODEL", "hermes3")
SCRIPT_TIMEOUT_SECONDS = int(os.getenv("SCRIPT_TIMEOUT_SECONDS", "60"))
MAX_REPAIR_ATTEMPTS = int(os.getenv("MAX_REPAIR_ATTEMPTS", "3"))
INVOKEAI_URL = os.getenv("INVOKEAI_URL", "http://localhost:9090").rstrip("/")
INVOKEAI_API_KEY = os.getenv("INVOKEAI_API_KEY", "")
INVOKEAI_WORKFLOW_PATH = os.getenv("INVOKEAI_WORKFLOW_PATH", "")
INVOKEAI_QUEUE_ID = os.getenv("INVOKEAI_QUEUE_ID", "default")
INVOKEAI_TIMEOUT_SECONDS = int(os.getenv("INVOKEAI_TIMEOUT_SECONDS", "240"))
INVOKEAI_POLL_SECONDS = float(os.getenv("INVOKEAI_POLL_SECONDS", "2"))
INVOKEAI_WIDTH = int(os.getenv("INVOKEAI_WIDTH", "1024"))
INVOKEAI_HEIGHT = int(os.getenv("INVOKEAI_HEIGHT", "1024"))
INVOKEAI_STEPS = int(os.getenv("INVOKEAI_STEPS", "30"))
INVOKEAI_CFG_SCALE = float(os.getenv("INVOKEAI_CFG_SCALE", "7.5"))
INVOKEAI_MODEL = os.getenv("INVOKEAI_MODEL", "")

client = AsyncOpenAI(base_url=LOCAL_LLM_URL, api_key="ollama")


CODE_BLOCK_RE = re.compile(r"```(?:python|py)?\s*\n(.*?)\n```", re.DOTALL | re.IGNORECASE)
URL_RE = re.compile(r"https?://[^\s<>)\"']+")
DOMAIN_RE = re.compile(r"\b(?:[a-z0-9-]+\.)+(?:com|org|net|edu|gov|io|co|uk|ca|au)\b", re.IGNORECASE)
USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
)
DEFAULT_NEWS_FEEDS = (
    "https://news.google.com/rss?hl=en-US&gl=US&ceid=US:en",
    "https://feeds.bbci.co.uk/news/rss.xml",
    "https://feeds.npr.org/1001/rss.xml",
)
WEB_PACKAGES = {
    "bs4": "beautifulsoup4",
    "requests": "requests",
}
SCRIPT_PACKAGES = {
    **WEB_PACKAGES,
    "fpdf": "fpdf",
    "matplotlib": "matplotlib",
    "pandas": "pandas",
    "docx": "python-docx",
    "PIL": "pillow",
    "reportlab": "reportlab",
}


def extract_python_code(text: str) -> str | None:
    match = CODE_BLOCK_RE.search(text or "")
    if match:
        return match.group(1).strip()

    # Some local models omit fences. Accept code-looking output as a fallback,
    # but do not mistake normal prose for an executable script.
    stripped = (text or "").strip()
    code_signals = ("import ", "from ", "def ", "class ", "print(", "with open(", "Path(")
    if "\n" in stripped and any(signal in stripped for signal in code_signals):
        return stripped
    return None


def choose_route(prompt: str) -> str:
    text = prompt.lower()

    image_words = (
        "generate image",
        "generate an image",
        "create image",
        "create an image",
        "make image",
        "make an image",
        "draw",
        "illustrate",
        "text to image",
        "txt2img",
        "image generation",
    )
    file_words = (
        "pdf",
        "docx",
        "word document",
        "txt",
        "text file",
        "html",
        "css",
        "javascript",
        "js",
        "spreadsheet",
        "xlsx",
        "csv",
        "chart",
        "graph",
        "plot",
        "generate a file",
        "create a file",
        "attachment",
        "downloadable",
        "save as",
    )
    document_words = (
        "pdf",
        "docx",
        "word document",
        "txt",
        "text file",
        "html",
        "document",
        "report",
        "attachment",
        "downloadable",
        "upload",
    )
    compute_words = (
        "calculate",
        "analyze this data",
        "run code",
        "script",
        "parse this",
        "convert",
        "extract",
    )
    live_web_words = (
        "latest",
        "today",
        "current",
        "recent",
        "news",
        "headline",
        "headlines",
        "web",
        "website",
        "search",
        "look up",
        "retrieve",
        "fetch",
        "live",
        "price",
        "weather",
        "stock",
    )

    has_file_request = any(word in text for word in file_words)
    has_web_request = any(word in text for word in live_web_words)
    has_document_request = any(word in text for word in document_words)
    has_image_request = any(word in text for word in image_words)

    if has_image_request:
        return "image"
    if has_web_request and has_document_request:
        return "web_document"
    if has_file_request:
        return "execute"
    if has_web_request:
        return "web"
    if any(word in text for word in compute_words):
        return "execute"
    return "chat"


async def ask_llm(messages, temperature=0.4) -> str:
    response = await client.chat.completions.create(
        model=LOCAL_LLM_MODEL,
        messages=messages,
        temperature=temperature,
    )
    return response.choices[0].message.content or ""


async def run_script(workspace_dir: Path, code: str) -> tuple[int, str, str]:
    script_path = workspace_dir / "generator.py"
    script_path.write_text(code, encoding="utf-8")

    proc = await asyncio.create_subprocess_exec(
        sys.executable,
        str(script_path),
        cwd=str(workspace_dir),
        stdout=asyncio.subprocess.PIPE,
        stderr=asyncio.subprocess.PIPE,
    )

    try:
        stdout, stderr = await asyncio.wait_for(
            proc.communicate(),
            timeout=SCRIPT_TIMEOUT_SECONDS,
        )
    except asyncio.TimeoutError:
        proc.kill()
        stdout, stderr = await proc.communicate()
        stderr += f"\nError: Script timed out after {SCRIPT_TIMEOUT_SECONDS} seconds.".encode()

    return (
        proc.returncode or 0,
        stdout.decode("utf-8", errors="replace").strip(),
        stderr.decode("utf-8", errors="replace").strip(),
    )


async def ensure_python_dependencies(workspace_dir: Path, required=None) -> None:
    required = required or SCRIPT_PACKAGES
    missing = []

    for import_name, package_name in required.items():
        proc = await asyncio.create_subprocess_exec(
            sys.executable,
            "-c",
            f"import {import_name}",
            cwd=str(workspace_dir),
            stdout=asyncio.subprocess.DEVNULL,
            stderr=asyncio.subprocess.DEVNULL,
        )
        await proc.communicate()
        if proc.returncode != 0:
            missing.append(package_name)

    if not missing:
        return

    print(f"Installing missing Python packages: {', '.join(missing)}")
    proc = await asyncio.create_subprocess_exec(
        sys.executable,
        "-m",
        "pip",
        "install",
        "-q",
        *missing,
        cwd=str(workspace_dir),
    )
    await proc.communicate()
    if proc.returncode != 0:
        raise RuntimeError(f"Failed to install Python packages: {', '.join(missing)}")


def requested_count(prompt: str, default: int = 3) -> int:
    match = re.search(r"\b(?:top|first|latest)\s+(\d{1,2})\b", prompt.lower())
    if not match:
        return default
    return max(1, min(int(match.group(1)), 10))


def safe_local_filename(filename: str, fallback: str = "attachment") -> str:
    name = Path(filename or fallback).name
    name = re.sub(r"[^a-zA-Z0-9._-]+", "_", name).strip("._")
    return name or fallback


def filename_with_mime_extension(filename: str, content_type: str) -> str:
    safe_name = safe_local_filename(filename)
    if Path(safe_name).suffix:
        return safe_name

    extension = mimetypes.guess_extension((content_type or "").split(";")[0].strip())
    if extension == ".jpe":
        extension = ".jpg"
    return f"{safe_name}{extension or ''}"


def unique_path(directory: Path, filename: str) -> Path:
    path = directory / safe_local_filename(filename)
    if not path.exists():
        return path

    stem = path.stem
    suffix = path.suffix
    for index in range(2, 1000):
        candidate = directory / f"{stem}_{index}{suffix}"
        if not candidate.exists():
            return candidate
    return directory / f"{stem}_{int(time.time())}{suffix}"


def save_incoming_attachments(workspace_dir: Path, exec_id: str, attachments: list[dict]) -> list[dict]:
    if not attachments:
        return []

    input_dir = workspace_dir / "inputs" / exec_id
    input_dir.mkdir(parents=True, exist_ok=True)
    saved = []

    for index, att in enumerate(attachments):
        raw_filename = att.get("filename") or f"attachment_{index + 1}"
        content_type = att.get("content_type") or mimetypes.guess_type(raw_filename)[0] or "application/octet-stream"
        filename = filename_with_mime_extension(raw_filename, content_type)
        record = {
            "filename": filename,
            "content_type": content_type,
            "path": "",
            "error": att.get("error"),
        }

        data = att.get("data")
        if not data:
            saved.append(record)
            continue

        path = unique_path(input_dir, filename)
        try:
            path.write_bytes(base64.b64decode(data))
            record["path"] = str(path)
            record["size"] = path.stat().st_size
        except Exception as exc:
            record["error"] = f"Failed to save attachment: {exc}"

        saved.append(record)

    return saved


def attachment_context(saved_attachments: list[dict]) -> str:
    if not saved_attachments:
        return ""

    lines = ["\n\nUser attached files available to this job:"]
    for index, att in enumerate(saved_attachments, start=1):
        path = att.get("path") or "(not saved)"
        filename = att.get("filename", f"attachment_{index}")
        content_type = att.get("content_type", "application/octet-stream")
        size = att.get("size")
        error = att.get("error")
        detail = f"{index}. {filename} ({content_type})"
        if size:
            detail += f", {size} bytes"
        detail += f"\n   Path: {path}"
        if error:
            detail += f"\n   Warning: {error}"
        lines.append(detail)

    lines.append(
        "When the user asks you to use an attached image or file as a reference, "
        "use these exact local paths. Do not invent attachment filenames."
    )
    return "\n".join(lines)


def invokeai_headers() -> dict[str, str]:
    headers = {"Accept": "application/json"}
    if INVOKEAI_API_KEY:
        headers["Authorization"] = f"Bearer {INVOKEAI_API_KEY}"
    return headers


def extract_image_prompt(prompt: str) -> str:
    text = prompt.strip()
    text = re.sub(r"<@!?\d+>", " ", text)
    patterns = (
        r"^\s*(?:please\s+)?(?:generate|create|make|draw|illustrate)\s+(?:me\s+)?(?:an?\s+)?image\s+(?:of|with|showing|for)?\s*",
        r"^\s*(?:please\s+)?(?:generate|create|make)\s+(?:me\s+)?(?:an?\s+)?(?:picture|art|illustration)\s+(?:of|with|showing|for)?\s*",
    )
    for pattern in patterns:
        text = re.sub(pattern, "", text, flags=re.IGNORECASE)
    return clean_text(text) or clean_text(prompt)


def inject_prompt_into_workflow(value, prompt: str):
    if isinstance(value, dict):
        updated = {}
        for key, child in value.items():
            lower_key = key.lower()
            if lower_key in {"prompt", "positive_prompt", "positiveprompt"} and isinstance(child, str):
                updated[key] = prompt
            elif lower_key in {"negative_prompt", "negativeprompt"} and isinstance(child, str):
                updated[key] = os.getenv("INVOKEAI_NEGATIVE_PROMPT", child)
            elif lower_key in {"width", "height", "steps"} and isinstance(child, int):
                if lower_key == "width":
                    updated[key] = INVOKEAI_WIDTH
                elif lower_key == "height":
                    updated[key] = INVOKEAI_HEIGHT
                else:
                    updated[key] = INVOKEAI_STEPS
            elif lower_key in {"cfg_scale", "cfgscale"} and isinstance(child, (int, float)):
                updated[key] = INVOKEAI_CFG_SCALE
            else:
                updated[key] = inject_prompt_into_workflow(child, prompt)
        return updated

    if isinstance(value, list):
        return [inject_prompt_into_workflow(item, prompt) for item in value]

    return value


def load_invokeai_workflow(prompt: str) -> dict | None:
    if not INVOKEAI_WORKFLOW_PATH:
        return None

    workflow_path = Path(INVOKEAI_WORKFLOW_PATH)
    workflow = json.loads(workflow_path.read_text(encoding="utf-8"))
    return inject_prompt_into_workflow(workflow, prompt)


def find_first_key(value, wanted_keys: set[str]):
    if isinstance(value, dict):
        for key, child in value.items():
            if key in wanted_keys and child:
                return child
            found = find_first_key(child, wanted_keys)
            if found:
                return found
    elif isinstance(value, list):
        for child in value:
            found = find_first_key(child, wanted_keys)
            if found:
                return found
    return None


def find_image_names(value) -> list[str]:
    names = []
    if isinstance(value, dict):
        for key, child in value.items():
            if key in {"image_name", "imageName", "image", "image_file", "imageFile", "name"} and isinstance(child, str):
                if re.search(r"\.(png|jpg|jpeg|webp)$", child, re.IGNORECASE) or key != "name":
                    names.append(child)
            elif key in {"images", "image_dtos", "imageDTOs"} and isinstance(child, list):
                names.extend(find_image_names(child))
            elif key in {"outputs", "results", "result"}:
                names.extend(find_image_names(child))
            else:
                names.extend(find_image_names(child))
    elif isinstance(value, list):
        for child in value:
            names.extend(find_image_names(child))
    return list(dict.fromkeys(names))


def find_session_id(value) -> str | None:
    if isinstance(value, dict):
        for key, child in value.items():
            if key in {"session_id", "sessionId"} and child:
                return str(child)
            found = find_session_id(child)
            if found:
                return found
    elif isinstance(value, list):
        for child in value:
            found = find_session_id(child)
            if found:
                return found
    return None


def fetch_invokeai_session_payload(session_id: str) -> dict | None:
    import requests

    candidates = [
        f"{INVOKEAI_URL}/api/v1/sessions/{session_id}",
        f"{INVOKEAI_URL}/api/v1/sessions/i/{session_id}",
    ]
    for url in candidates:
        try:
            response = requests.get(url, headers=invokeai_headers(), timeout=30)
            response.raise_for_status()
            return response.json()
        except Exception:
            continue
    return None


def save_image_bytes(workspace_dir: Path, image_bytes: bytes, filename: str = "invokeai_image.png") -> Path:
    images_dir = workspace_dir / "generated_images"
    images_dir.mkdir(parents=True, exist_ok=True)
    path = unique_path(images_dir, filename)
    path.write_bytes(image_bytes)
    return path


def download_invokeai_image(image_name: str) -> bytes:
    import requests

    candidates = [
        f"{INVOKEAI_URL}/api/v1/images/i/{image_name}/full",
        f"{INVOKEAI_URL}/api/v1/images/{image_name}/full",
        f"{INVOKEAI_URL}/outputs/images/{image_name}",
    ]
    last_error = None
    for url in candidates:
        try:
            response = requests.get(url, headers=invokeai_headers(), timeout=60)
            response.raise_for_status()
            return response.content
        except Exception as exc:
            last_error = exc
    raise RuntimeError(f"Could not download InvokeAI image {image_name}: {last_error}")


def save_image_from_response(payload, workspace_dir: Path) -> Path | None:
    if isinstance(payload, dict):
        base64_data = find_first_key(payload, {"b64_json", "base64", "image_base64", "data"})
        if isinstance(base64_data, str) and len(base64_data) > 100:
            if "," in base64_data and base64_data.split(",", 1)[0].startswith("data:"):
                base64_data = base64_data.split(",", 1)[1]
            return save_image_bytes(workspace_dir, base64.b64decode(base64_data), "invokeai_image.png")

        image_url = find_first_key(payload, {"url", "image_url", "imageUrl"})
        if isinstance(image_url, str) and image_url.startswith("http"):
            import requests

            response = requests.get(image_url, headers=invokeai_headers(), timeout=60)
            response.raise_for_status()
            extension = Path(urlparse(image_url).path).suffix or ".png"
            return save_image_bytes(workspace_dir, response.content, f"invokeai_image{extension}")

        image_names = find_image_names(payload)
        if image_names:
            image_name = image_names[0]
            return save_image_bytes(workspace_dir, download_invokeai_image(image_name), Path(image_name).name)

    return None


def enqueue_invokeai_workflow(workflow: dict) -> dict:
    import requests

    enqueue_url = f"{INVOKEAI_URL}/api/v1/queue/{INVOKEAI_QUEUE_ID}/enqueue_batch"
    payload_candidates = [
        {"batch": {"graph": workflow, "runs": 1}, "prepend": False},
        {"graph": workflow, "runs": 1, "prepend": False},
        workflow,
    ]
    last_error = None
    for payload in payload_candidates:
        try:
            response = requests.post(enqueue_url, json=payload, headers=invokeai_headers(), timeout=30)
            response.raise_for_status()
            return response.json()
        except Exception as exc:
            last_error = exc
    raise RuntimeError(f"InvokeAI queue request failed: {last_error}")


def poll_invokeai_queue(enqueue_response: dict) -> dict:
    import requests

    item_id = find_first_key(enqueue_response, {"item_id", "itemId", "id"})
    if not item_id:
        return enqueue_response

    deadline = time.time() + INVOKEAI_TIMEOUT_SECONDS
    last_payload = enqueue_response
    status_url = f"{INVOKEAI_URL}/api/v1/queue/{INVOKEAI_QUEUE_ID}/i/{item_id}"

    while time.time() < deadline:
        response = requests.get(status_url, headers=invokeai_headers(), timeout=30)
        response.raise_for_status()
        last_payload = response.json()

        status = str(find_first_key(last_payload, {"status", "state"}) or "").lower()
        if status in {"completed", "complete", "succeeded", "success", "failed", "canceled", "cancelled"}:
            return last_payload
        if find_image_names(last_payload):
            return last_payload
        time.sleep(INVOKEAI_POLL_SECONDS)

    raise TimeoutError(f"InvokeAI generation did not finish within {INVOKEAI_TIMEOUT_SECONDS} seconds")


def invokeai_simple_generate(prompt: str, workspace_dir: Path) -> Path:
    import requests

    payload = {
        "prompt": prompt,
        "width": INVOKEAI_WIDTH,
        "height": INVOKEAI_HEIGHT,
        "steps": INVOKEAI_STEPS,
        "cfg_scale": INVOKEAI_CFG_SCALE,
    }
    if INVOKEAI_MODEL:
        payload["model"] = INVOKEAI_MODEL

    endpoints = [
        f"{INVOKEAI_URL}/api/v1/images/generations",
        f"{INVOKEAI_URL}/api/generate",
    ]
    errors = []
    for endpoint in endpoints:
        try:
            response = requests.post(endpoint, json=payload, headers=invokeai_headers(), timeout=INVOKEAI_TIMEOUT_SECONDS)
            response.raise_for_status()
            image_path = save_image_from_response(response.json(), workspace_dir)
            if image_path:
                return image_path
            errors.append(f"{endpoint}: response did not contain an image")
        except Exception as exc:
            errors.append(f"{endpoint}: {exc}")

    raise RuntimeError(
        "InvokeAI did not accept a simple generation request. Set INVOKEAI_WORKFLOW_PATH "
        "to an exported InvokeAI workflow JSON template. Attempts:\n" + "\n".join(errors)
    )


def generate_image_with_invokeai_sync(prompt: str, workspace_dir: Path) -> Path:
    image_prompt = extract_image_prompt(prompt)
    workflow = load_invokeai_workflow(image_prompt)

    if workflow:
        enqueue_response = enqueue_invokeai_workflow(workflow)
        final_payload = poll_invokeai_queue(enqueue_response)
        image_path = save_image_from_response(final_payload, workspace_dir)
        if image_path:
            return image_path
        session_id = find_session_id(final_payload) or find_session_id(enqueue_response)
        if session_id:
            session_payload = fetch_invokeai_session_payload(session_id)
            image_path = save_image_from_response(session_payload, workspace_dir)
            if image_path:
                return image_path
        raise RuntimeError("InvokeAI completed, but no generated image name or bytes were found in the response.")

    return invokeai_simple_generate(image_prompt, workspace_dir)


async def generate_image_with_invokeai(prompt: str, workspace_dir: Path) -> str:
    await ensure_python_dependencies(workspace_dir, WEB_PACKAGES)
    image_path = await asyncio.to_thread(generate_image_with_invokeai_sync, prompt, workspace_dir)
    return f"Generated image from InvokeAI: `{image_path.name}`"


def clean_text(value: str | None) -> str:
    if not value:
        return ""
    value = unescape(value)
    value = re.sub(r"\s+", " ", value).strip()
    return value


def normalize_url(raw_url: str) -> str:
    return raw_url.rstrip(".,);]}")


def prompt_urls(prompt: str) -> list[str]:
    urls = [normalize_url(match.group(0)) for match in URL_RE.finditer(prompt)]
    domains = [
        match.group(0).lower()
        for match in DOMAIN_RE.finditer(prompt)
        if not match.group(0).lower().startswith(("http.", "https."))
    ]
    for domain in domains:
        url = f"https://{domain}"
        if url not in urls:
            urls.append(url)
    return urls


def search_query_from_prompt(prompt: str) -> str:
    query = URL_RE.sub(" ", prompt)
    query = DOMAIN_RE.sub(" ", query)
    replacements = (
        "search the web for",
        "search the web",
        "search web for",
        "search web",
        "look up",
        "retrieve",
        "fetch",
        "get me",
        "get",
        "tell me about",
        "what is happening with",
        "what's happening with",
        "top",
        "first",
        "make a pdf",
        "make",
        "create a pdf",
        "create",
        "put it in a pdf",
        "into a pdf",
        "pdf",
        "docx",
        "word document",
        "html",
        "txt",
        "text file",
        "file",
        "report",
        "document",
        "them",
        "it",
        "upload",
        "easy to read",
        "sources",
        "source",
        "links",
        "link",
        "headlines",
        "headline",
        "summary",
        "summarize",
        "summarise",
        "digest",
        "latest",
        "current",
        "recent",
        "news",
        "today",
        "please",
    )
    lowered = query.lower()
    for phrase in replacements:
        pattern = r"\b" + re.escape(phrase).replace(r"\ ", r"\s+") + r"\b"
        lowered = re.sub(pattern, " ", lowered)
    lowered = re.sub(r"\b(?:the|a|an|and|to|for|with|of|on|in|into|about|me)\b", " ", lowered)
    lowered = re.sub(r"\b\d{1,2}\b", " ", lowered)
    return re.sub(r"\s+", " ", lowered).strip(" ?.,")


def wants_sources(prompt: str) -> bool:
    text = prompt.lower()
    return any(word in text for word in ("source", "sources", "link", "links", "url", "urls", "citation"))


def wants_summary(prompt: str) -> bool:
    text = prompt.lower()
    return any(word in text for word in ("summary", "summarize", "summarise", "digest", "briefing", "recap"))


def is_general_news_request(prompt: str) -> bool:
    text = prompt.lower()
    if prompt_urls(prompt):
        return False
    return "news" in text and not search_query_from_prompt(prompt)


def requested_document_format(prompt: str) -> str:
    text = prompt.lower()
    if "docx" in text or "word document" in text:
        return "docx"
    if "html" in text or "web page" in text:
        return "html"
    if "txt" in text or "text file" in text:
        return "txt"
    return "pdf"


def source_name(url: str) -> str:
    host = urlparse(url).netloc.lower().removeprefix("www.")
    return host or url


def parse_feed_items(feed_text: str, feed_url: str) -> list[dict[str, str]]:
    try:
        root = ET.fromstring(feed_text.encode("utf-8"))
    except ET.ParseError:
        return []

    items = []
    for item in root.findall(".//item"):
        title = clean_text(item.findtext("title"))
        link = clean_text(item.findtext("link"))
        description = clean_text(item.findtext("description"))
        published = clean_text(item.findtext("pubDate"))
        if title:
            items.append(
                {
                    "title": title,
                    "url": link or feed_url,
                    "source": source_name(feed_url),
                    "snippet": description,
                    "published": published,
                }
            )

    if items:
        return items

    namespace = {"atom": "http://www.w3.org/2005/Atom"}
    for entry in root.findall(".//atom:entry", namespace):
        title = clean_text(entry.findtext("atom:title", default="", namespaces=namespace))
        link = ""
        for link_node in entry.findall("atom:link", namespace):
            href = link_node.attrib.get("href")
            if href:
                link = href
                break
        if title:
            items.append(
                {
                    "title": title,
                    "url": link or feed_url,
                    "source": source_name(feed_url),
                    "snippet": "",
                    "published": "",
                }
            )

    return items


def fetch_text(url: str) -> tuple[str, str]:
    import requests

    response = requests.get(
        url,
        headers={"User-Agent": USER_AGENT, "Accept": "text/html,application/rss+xml,application/xml"},
        timeout=12,
    )
    response.raise_for_status()
    return response.text, response.url


def extract_article_text(url: str) -> str:
    from bs4 import BeautifulSoup

    html_text, final_url = fetch_text(url)
    soup = BeautifulSoup(html_text, "html.parser")
    for noisy in soup(["script", "style", "nav", "footer", "aside", "form", "noscript"]):
        noisy.decompose()

    article = soup.find("article") or soup.find("main") or soup.body or soup
    paragraphs = []
    for node in article.find_all(["p", "li"]):
        text = clean_text(node.get_text(" ", strip=True))
        if 60 <= len(text) <= 1200:
            paragraphs.append(text)

    seen = set()
    cleaned = []
    for paragraph in paragraphs:
        key = paragraph.lower()
        if key in seen:
            continue
        cleaned.append(paragraph)
        seen.add(key)
        if sum(len(item) for item in cleaned) > 7000:
            break

    return "\n\n".join(cleaned)


def discover_feed_urls(page_url: str, html_text: str) -> list[str]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html_text, "html.parser")
    feed_urls = []

    for node in soup.find_all("link"):
        type_value = (node.get("type") or "").lower()
        rel_value = " ".join(node.get("rel") or []).lower()
        href = node.get("href")
        if href and ("rss" in type_value or "atom" in type_value or "alternate" in rel_value):
            feed_urls.append(urljoin(page_url, href))

    parsed = urlparse(page_url)
    base = f"{parsed.scheme}://{parsed.netloc}"
    feed_urls.extend(
        [
            urljoin(base, "/feed"),
            urljoin(base, "/rss"),
            urljoin(base, "/rss.xml"),
            urljoin(base, "/feed.xml"),
        ]
    )

    return list(dict.fromkeys(feed_urls))


def extract_headlines_from_html(page_url: str, html_text: str) -> list[dict[str, str]]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html_text, "html.parser")
    for noisy in soup(["script", "style", "nav", "footer", "aside", "form"]):
        noisy.decompose()

    selectors = (
        "article h1 a",
        "article h2 a",
        "article h3 a",
        "main h1 a",
        "main h2 a",
        "main h3 a",
        "h1 a",
        "h2 a",
        "h3 a",
        "article h1",
        "article h2",
        "article h3",
        "main h1",
        "main h2",
        "main h3",
        "h1",
        "h2",
    )
    skip_words = {
        "subscribe",
        "sign in",
        "log in",
        "advertisement",
        "privacy policy",
        "terms of use",
        "newsletter",
        "more news",
    }

    items = []
    seen = set()
    for selector in selectors:
        for node in soup.select(selector):
            title = clean_text(node.get_text(" ", strip=True))
            if not title or len(title) < 12 or len(title) > 220:
                continue
            if title.lower() in skip_words or title in seen:
                continue

            link_node = node if node.name == "a" else node.find("a")
            href = link_node.get("href") if link_node else None
            items.append(
                {
                    "title": title,
                    "url": urljoin(page_url, href) if href else page_url,
                    "source": source_name(page_url),
                    "snippet": "",
                    "published": "",
                }
            )
            seen.add(title)

    return items


def search_duckduckgo(query: str, limit: int) -> list[dict[str, str]]:
    import requests
    from bs4 import BeautifulSoup

    url = f"https://html.duckduckgo.com/html/?q={quote_plus(query)}"
    response = requests.get(url, headers={"User-Agent": USER_AGENT}, timeout=12)
    response.raise_for_status()

    soup = BeautifulSoup(response.text, "html.parser")
    results = []
    for result in soup.select(".result"):
        link_node = result.select_one(".result__a")
        if not link_node:
            continue
        title = clean_text(link_node.get_text(" ", strip=True))
        href = link_node.get("href") or ""
        snippet_node = result.select_one(".result__snippet")
        snippet = clean_text(snippet_node.get_text(" ", strip=True) if snippet_node else "")
        if title and href:
            results.append(
                {
                    "title": title,
                    "url": href,
                    "source": source_name(href),
                    "snippet": snippet,
                    "published": "",
                }
            )
        if len(results) >= limit:
            break
    return results


def search_live_web_sync(prompt: str, limit: int = 5) -> list[dict[str, str]]:
    query = search_query_from_prompt(prompt) or "top news"
    results = []

    news_feed = f"https://news.google.com/rss/search?q={quote_plus(query)}&hl=en-US&gl=US&ceid=US:en"
    try:
        feed_text, final_feed_url = fetch_text(news_feed)
        results.extend(parse_feed_items(feed_text, final_feed_url))
    except Exception:
        pass

    if len(results) < limit:
        try:
            results.extend(search_duckduckgo(query, limit))
        except Exception:
            pass

    deduped = []
    seen = set()
    for item in results:
        key = (item.get("title") or "").lower()
        if not key or key in seen:
            continue
        deduped.append(item)
        seen.add(key)
        if len(deduped) >= limit:
            break

    return deduped


def fetch_prompt_url_items(prompt: str, limit: int) -> list[dict[str, str]]:
    items = []
    for url in prompt_urls(prompt):
        text, final_url = fetch_text(url)
        feed_items = parse_feed_items(text, final_url)
        if feed_items:
            items.extend(feed_items)
        else:
            for feed_url in discover_feed_urls(final_url, text):
                try:
                    feed_text, final_feed_url = fetch_text(feed_url)
                    feed_items = parse_feed_items(feed_text, final_feed_url)
                    if feed_items:
                        items.extend(feed_items)
                        break
                except Exception:
                    continue
            if not items:
                items.extend(extract_headlines_from_html(final_url, text))
        if len(items) >= limit:
            break
    return items[:limit]


def format_headlines(items: list[dict[str, str]], limit: int) -> str:
    if not items:
        return "I could not find headlines from that source."

    lines = ["**Top headlines:**"]
    for index, item in enumerate(items[:limit], start=1):
        title = item["title"]
        url = item.get("url") or ""
        source = item.get("source") or source_name(url)
        if url:
            lines.append(f"{index}. {title}\n   Source: {source}\n   {url}")
        else:
            lines.append(f"{index}. {title}\n   Source: {source}")
    return "\n".join(lines)


def dedupe_items(items: list[dict[str, str]], limit: int) -> list[dict[str, str]]:
    deduped = []
    seen = set()
    for item in items:
        title = item.get("title", "")
        key = title.lower()
        if not key or key in seen:
            continue
        deduped.append(item)
        seen.add(key)
        if len(deduped) >= limit:
            break
    return deduped


def collect_headline_items_sync(prompt: str, limit: int) -> tuple[list[dict[str, str]], list[str]]:
    urls = prompt_urls(prompt)
    items = []
    errors = []

    if urls:
        try:
            return dedupe_items(fetch_prompt_url_items(prompt, limit), limit), errors
        except Exception as exc:
            errors.append(str(exc))
            return [], errors

    candidates = list(DEFAULT_NEWS_FEEDS)
    for url in candidates:
        try:
            text, final_url = fetch_text(url)
            items.extend(parse_feed_items(text, final_url))
        except Exception as exc:
            errors.append(f"{source_name(url)}: {exc}")

        if len(items) >= limit:
            break

    return dedupe_items(items, limit), errors


def retrieve_headlines_sync(prompt: str, limit: int) -> str:
    deduped, errors = collect_headline_items_sync(prompt, limit)
    if deduped:
        return format_headlines(deduped, limit)

    if errors:
        return "I could not retrieve headlines.\n\n" + "\n".join(f"- {error}" for error in errors[:3])
    return "I could not find headlines from that source."


async def retrieve_headlines(prompt: str, workspace_dir: Path) -> str:
    await ensure_python_dependencies(workspace_dir, WEB_PACKAGES)
    limit = requested_count(prompt)
    return await asyncio.to_thread(retrieve_headlines_sync, prompt, limit)


def format_web_results(items: list[dict[str, str]], limit: int) -> str:
    if not items:
        return "I could not find live web results for that request."

    lines = ["**Live web results:**"]
    for index, item in enumerate(items[:limit], start=1):
        title = item.get("title", "Untitled")
        snippet = item.get("snippet", "")
        published = item.get("published", "")
        url = item.get("url", "")
        source = item.get("source") or source_name(url)
        lines.append(f"{index}. {title}")
        if snippet:
            lines.append(f"   {snippet}")
        if published:
            lines.append(f"   Published: {published}")
        if source:
            lines.append(f"   Source: {source}")
        if url:
            lines.append(f"   {url}")
    return "\n".join(lines)


def enrich_items_with_article_text_sync(items: list[dict[str, str]]) -> list[dict[str, str]]:
    enriched = []
    for item in items:
        copy = dict(item)
        article_text = ""
        url = copy.get("url", "")
        if url:
            try:
                article_text = extract_article_text(url)
            except Exception:
                article_text = ""
        copy["article_text"] = article_text or copy.get("snippet", "")
        enriched.append(copy)
    return enriched


async def answer_from_live_web(prompt: str, workspace_dir: Path) -> str:
    await ensure_python_dependencies(workspace_dir, WEB_PACKAGES)

    if "headline" in prompt.lower() or "headlines" in prompt.lower():
        return await retrieve_headlines(prompt, workspace_dir)

    if is_general_news_request(prompt):
        items, _errors = await asyncio.to_thread(collect_headline_items_sync, prompt, 5)
    else:
        items = await asyncio.to_thread(search_live_web_sync, prompt, 5)

    if not items:
        return "I could not find live web results for that request."

    if wants_summary(prompt):
        enriched_items = await asyncio.to_thread(enrich_items_with_article_text_sync, items)
        return await build_news_digest(prompt, enriched_items)

    source_context = "\n\n".join(
        (
            f"Title: {item.get('title', '')}\n"
            f"Snippet: {item.get('snippet', '')}\n"
            f"Published: {item.get('published', '')}\n"
            f"Source: {item.get('source', '')}\n"
            f"URL: {item.get('url', '')}"
        )
        for item in items
    )
    system_prompt = """You answer using only the live web result snippets provided.

If the snippets are thin or incomplete, say so. Include source links in the
answer. Do not add facts that are not supported by the provided results."""
    summary = await ask_llm(
        [
            {"role": "system", "content": system_prompt},
            {
                "role": "user",
                "content": f"User request: {prompt}\n\nLive web results:\n{source_context}",
            },
        ],
        temperature=0.2,
    )
    return summary.strip() or format_web_results(items, 5)


def safe_filename(value: str, fallback: str = "web_report") -> str:
    value = re.sub(r"[^a-zA-Z0-9_-]+", "_", value.lower()).strip("_")
    return value[:40] or fallback


def report_basename(prompt: str, source_mode: bool) -> str:
    if "headline" in prompt.lower() or "headlines" in prompt.lower() or is_general_news_request(prompt):
        return "headlines_sources" if source_mode else "news_digest"
    return safe_filename(search_query_from_prompt(prompt), "web_report")


def plain_digest_from_items(prompt: str, items: list[dict[str, str]]) -> str:
    lines = ["Hermes News Digest", "", clean_text(prompt), ""]
    for index, item in enumerate(items, start=1):
        title = clean_text(item.get("title", "Untitled"))
        source = clean_text(item.get("source", ""))
        body = clean_text(item.get("article_text") or item.get("snippet") or "")
        lines.append(f"{index}. {title}")
        if source:
            lines.append(f"Source: {source}")
        if body:
            lines.append(body[:1400])
        lines.append("")
    return "\n".join(lines).strip()


async def build_news_digest(prompt: str, items: list[dict[str, str]]) -> str:
    source_context = "\n\n".join(
        (
            f"Topic {index}: {item.get('title', '')}\n"
            f"Source: {item.get('source', '')}\n"
            f"Published: {item.get('published', '')}\n"
            f"Article text or snippet:\n{(item.get('article_text') or item.get('snippet') or '')[:5000]}"
        )
        for index, item in enumerate(items, start=1)
    )
    system_prompt = """Write an easy-to-read news digest from the provided article text.

Use a short title, a brief overview paragraph, and clear sections for each major
topic. Explain what happened, why it matters, and any uncertainty or limits in
the available reporting. Do not include raw URLs unless the user requested
sources or links. Do not invent facts beyond the provided article text."""
    digest = await ask_llm(
        [
            {"role": "system", "content": system_prompt},
            {
                "role": "user",
                "content": f"User request: {prompt}\n\nFetched articles:\n{source_context}",
            },
        ],
        temperature=0.25,
    )
    return digest.strip() or plain_digest_from_items(prompt, items)


def markdownish_to_blocks(text: str) -> list[tuple[str, str]]:
    blocks = []
    for raw_line in text.splitlines():
        line = clean_text(raw_line.strip("#* "))
        if not line:
            blocks.append(("spacer", ""))
        elif raw_line.startswith("#") or len(line) < 90 and not line.endswith("."):
            blocks.append(("heading", line))
        else:
            blocks.append(("paragraph", line))
    return blocks


def create_pdf_report(
    prompt: str,
    workspace_dir: Path,
    digest_text: str | None = None,
    items: list[dict[str, str]] | None = None,
    source_mode: bool = False,
) -> str:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    filename = f"{report_basename(prompt, source_mode)}.pdf"
    path = workspace_dir / filename
    styles = getSampleStyleSheet()
    title = "Hermes Source Report" if source_mode else "Hermes News Digest"
    doc = SimpleDocTemplate(str(path), pagesize=letter, title=title)
    story = []

    story.append(Paragraph(title, styles["Title"]))
    story.append(Paragraph(clean_text(prompt), styles["Normal"]))
    story.append(Spacer(1, 12))

    if source_mode:
        for index, item in enumerate(items or [], start=1):
            item_title = clean_text(item.get("title", "Untitled"))
            snippet = clean_text(item.get("snippet", ""))
            published = clean_text(item.get("published", ""))
            source = clean_text(item.get("source", ""))
            url = clean_text(item.get("url", ""))

            story.append(Paragraph(f"{index}. {item_title}", styles["Heading2"]))
            if snippet:
                story.append(Paragraph(snippet, styles["BodyText"]))
            if published:
                story.append(Paragraph(f"Published: {published}", styles["BodyText"]))
            if source:
                story.append(Paragraph(f"Source: {source}", styles["BodyText"]))
            if url:
                story.append(Paragraph(url, styles["BodyText"]))
            story.append(Spacer(1, 10))
    else:
        for block_type, value in markdownish_to_blocks(digest_text or ""):
            if block_type == "spacer":
                story.append(Spacer(1, 8))
            elif block_type == "heading":
                story.append(Paragraph(value, styles["Heading2"]))
            else:
                story.append(Paragraph(value, styles["BodyText"]))

    doc.build(story)
    return filename


def create_txt_report(prompt: str, workspace_dir: Path, digest_text: str, source_mode: bool) -> str:
    filename = f"{report_basename(prompt, source_mode)}.txt"
    path = workspace_dir / filename
    path.write_text(digest_text, encoding="utf-8")
    return filename


def create_html_report(prompt: str, workspace_dir: Path, digest_text: str, source_mode: bool) -> str:
    filename = f"{report_basename(prompt, source_mode)}.html"
    title = "Hermes Source Report" if source_mode else "Hermes News Digest"
    body_parts = [f"<h1>{escape(title)}</h1>", f"<p>{escape(clean_text(prompt))}</p>"]
    for block_type, value in markdownish_to_blocks(digest_text):
        if block_type == "spacer":
            continue
        tag = "h2" if block_type == "heading" else "p"
        body_parts.append(f"<{tag}>{escape(value)}</{tag}>")
    html = (
        "<!doctype html><html><head><meta charset=\"utf-8\">"
        f"<title>{escape(title)}</title>"
        "<style>body{font-family:Arial,sans-serif;line-height:1.55;max-width:820px;margin:40px auto;padding:0 24px}"
        "h1,h2{line-height:1.2}p{margin:0 0 14px}</style></head><body>"
        + "\n".join(body_parts)
        + "</body></html>"
    )
    path = workspace_dir / filename
    path.write_text(html, encoding="utf-8")
    return filename


def create_docx_report(prompt: str, workspace_dir: Path, digest_text: str, source_mode: bool) -> str:
    from docx import Document

    filename = f"{report_basename(prompt, source_mode)}.docx"
    title = "Hermes Source Report" if source_mode else "Hermes News Digest"
    document = Document()
    document.add_heading(title, 0)
    document.add_paragraph(clean_text(prompt))

    for block_type, value in markdownish_to_blocks(digest_text):
        if block_type == "spacer":
            continue
        if block_type == "heading":
            document.add_heading(value, level=1)
        else:
            document.add_paragraph(value)

    document.save(str(workspace_dir / filename))
    return filename


def create_document_report(
    prompt: str,
    workspace_dir: Path,
    fmt: str,
    digest_text: str,
    items: list[dict[str, str]],
    source_mode: bool,
) -> str:
    if source_mode:
        digest_text = format_web_results(items, len(items))

    if fmt == "docx":
        return create_docx_report(prompt, workspace_dir, digest_text, source_mode)
    if fmt == "html":
        return create_html_report(prompt, workspace_dir, digest_text, source_mode)
    if fmt == "txt":
        return create_txt_report(prompt, workspace_dir, digest_text, source_mode)
    return create_pdf_report(prompt, workspace_dir, digest_text, items, source_mode)


def snapshot_workspace_files(workspace_dir: Path) -> set[str]:
    files = set()
    for path in workspace_dir.rglob("*"):
        if path.is_file():
            files.add(str(path.relative_to(workspace_dir)))
    return files


async def create_web_document(prompt: str, workspace_dir: Path) -> str:
    await ensure_python_dependencies(workspace_dir, SCRIPT_PACKAGES)
    limit = requested_count(prompt, default=5)
    fmt = requested_document_format(prompt)
    source_mode = wants_sources(prompt)

    if "headline" in prompt.lower() or "headlines" in prompt.lower() or is_general_news_request(prompt):
        items, _errors = await asyncio.to_thread(collect_headline_items_sync, prompt, limit)
    else:
        items = await asyncio.to_thread(search_live_web_sync, prompt, limit)

    if not items:
        return f"I could not find live web results to put into a {fmt.upper()} file."

    if source_mode:
        digest_text = format_web_results(items, len(items))
    else:
        enriched_items = await asyncio.to_thread(enrich_items_with_article_text_sync, items)
        digest_text = await build_news_digest(prompt, enriched_items)
        items = enriched_items

    filename = await asyncio.to_thread(
        create_document_report,
        prompt,
        workspace_dir,
        fmt,
        digest_text,
        items,
        source_mode,
    )
    mode_label = "source list" if source_mode else "summarized digest"
    return f"Created an easy-to-read {fmt.upper()} {mode_label} with {len(items)} live web result(s): `{filename}`"


async def direct_answer(prompt: str) -> str:
    system_prompt = """You are Hermes, a helpful agent inside Fluxer.

Answer normal questions directly. Do not claim you used tools when you did not.
If the user asks for current/live information that you cannot know, say that a
tool run is needed instead of guessing. Keep answers concise and useful."""

    return await ask_llm(
        [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": prompt},
        ],
        temperature=0.5,
    )


async def execute_with_python(prompt: str, workspace_dir: Path) -> str:
    system_prompt = f"""You are Hermes, an agent that may write and run Python when a task needs tools.

Use Python only for tasks that need file generation, calculation, data parsing,
charts, PDFs, DOCX files, TXT, HTML, CSS, JS, spreadsheets, or live/current web
information. The final answer will come from the script stdout and generated files.

Workspace path: {workspace_dir}

Rules for scripts:
1. Output exactly one complete Python script in a fenced ```python block.
2. Always print a useful final result to stdout.
3. Save any generated user files into the workspace directory. You may create
   common attachment types such as PDF, DOCX, TXT, HTML, CSS, JS, JSON, CSV,
   XLSX, images, archives, and source code when the user asks for them.
   If the user attached files, their exact local paths are listed in the user
   message. Use those paths directly for reading, embedding, or reference.
   For PDF generation, prefer reportlab unless the user asks for another
   library. If you use Pillow images with BytesIO, always pass an explicit
   format, for example image.save(buffer, format="PNG"). If you save an image
   to disk, always use a filename with a real extension such as .png or .jpg.
4. For web/current-information tasks, use requests with a realistic User-Agent,
   short timeouts, and graceful failure messages. Do not fabricate results.
5. Handle empty results and missing fields without crashing.
6. Do not use infinite loops or interactive input."""

    conversation = [
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": prompt},
    ]

    await ensure_python_dependencies(workspace_dir, SCRIPT_PACKAGES)

    last_error = ""
    for attempt in range(1, MAX_REPAIR_ATTEMPTS + 1):
        print(f"Asking Ollama for executable plan (attempt {attempt}/{MAX_REPAIR_ATTEMPTS})...")
        ai_text = await ask_llm(conversation, temperature=0.2)
        conversation.append({"role": "assistant", "content": ai_text})

        code = extract_python_code(ai_text)
        if not code:
            last_error = "The model did not return a Python script."
            conversation.append(
                {
                    "role": "user",
                    "content": (
                        "Return one complete Python script wrapped in ```python fences. "
                        "The script must print the final result."
                    ),
                }
            )
            continue

        returncode, stdout, stderr = await run_script(workspace_dir, code)
        if returncode == 0 and stdout:
            return f"**Execution Output:**\n\n{stdout}"

        if returncode == 0:
            last_error = "The script ran successfully but printed nothing."
            conversation.append(
                {
                    "role": "user",
                    "content": (
                        "The script ran successfully but printed nothing. Rewrite it so "
                        "it prints the useful final result to stdout."
                    ),
                }
            )
        else:
            last_error = stderr or f"Script exited with code {returncode}."
            repair_hint = ""
            lowered_error = last_error.lower()
            if "unknown file extension" in lowered_error and "pil" in lowered_error:
                repair_hint = (
                    "\n\nRepair hint: Pillow cannot infer an image format when saving to "
                    "a BytesIO object or extensionless path. Use image.save(buffer, "
                    "format=\"PNG\") for BytesIO, or save to a path ending in .png/.jpg. "
                    "For PDFs, reportlab can embed the original attached image path "
                    "directly with drawImage/ImageReader."
                )
            conversation.append(
                {
                    "role": "user",
                    "content": (
                        "The script failed with this error:\n"
                        f"{last_error}\n\n"
                        "Fix the bug and return the complete Python script again."
                        f"{repair_hint}"
                    ),
                }
            )

    return f"I tried to run the tool workflow, but it did not complete cleanly:\n\n```text\n{last_error}\n```"


def collect_attachments(workspace_dir: Path, initial_files: set[str]) -> list[dict[str, str]]:
    attachments = []
    generated_files = snapshot_workspace_files(workspace_dir) - initial_files

    for relative_name in sorted(generated_files):
        path = workspace_dir / relative_name
        if (
            not path.is_file()
            or path.name == "generator.py"
            or "__pycache__" in path.parts
            or path.suffix in {".pyc", ".pyo"}
        ):
            continue
        with path.open("rb") as file:
            attachments.append(
                {
                    "filename": path.name,
                    "data": base64.b64encode(file.read()).decode("utf-8"),
                }
            )

    return attachments


async def run_worker():
    nc = NATS()
    await nc.connect(NATS_URL)
    r = redis.from_url(REDIS_URL, decode_responses=True)

    print("Hermes worker connected to NATS and Redis")
    print(f"Pointing to Ollama-compatible API at: {LOCAL_LLM_URL}")

    async def message_handler(msg):
        data = json.loads(msg.data.decode())
        exec_id = data["metadata"]["execution_id"]
        channel_id = data["payload"]["channel_id"]
        user_id = data["payload"]["user_id"]
        original_prompt = data["payload"]["prompt"]
        incoming_attachments = data["payload"].get("attachments", [])

        print(f"\nPicked up job {exec_id}. Prompt: {original_prompt!r}")

        lock_key = f"lock:job:{exec_id}"
        if not await r.set(lock_key, "true", nx=True, ex=300):
            print(f"Ignoring duplicate job delivery for {exec_id}")
            return

        workspace_dir = Path(os.getcwd()) / "workspaces" / str(user_id)
        workspace_dir.mkdir(parents=True, exist_ok=True)
        saved_attachments = save_incoming_attachments(workspace_dir, exec_id, incoming_attachments)
        if saved_attachments:
            print(f"Saved {len(saved_attachments)} incoming attachment(s) for {exec_id}")
        initial_files = snapshot_workspace_files(workspace_dir)
        prompt = original_prompt + attachment_context(saved_attachments)

        async def update_state(text, is_final=False, attachments=None):
            await nc.publish(
                "hermes.execution.state_change",
                json.dumps(
                    {
                        "event": "state_transition",
                        "metadata": {"execution_id": exec_id},
                        "payload": {
                            "channel_id": channel_id,
                            "display_text": text,
                            "is_final": is_final,
                            "attachments": attachments or [],
                        },
                    }
                ).encode(),
            )

        try:
            route = choose_route(original_prompt)
            await update_state("Thinking..." if route == "chat" else "Working on it...")

            if route == "chat":
                final_answer = await direct_answer(prompt)
            elif route == "image":
                final_answer = await generate_image_with_invokeai(prompt, workspace_dir)
            elif route == "web":
                final_answer = await answer_from_live_web(prompt, workspace_dir)
            elif route == "web_document":
                final_answer = await create_web_document(prompt, workspace_dir)
            else:
                final_answer = await execute_with_python(prompt, workspace_dir)

            if not final_answer.strip():
                final_answer = "Finished processing."

            attachments = collect_attachments(workspace_dir, initial_files)
            await update_state(final_answer, is_final=True, attachments=attachments)
            print(f"Finished execution {exec_id}")

        except Exception as exc:
            print(f"Fatal error: {exc}")
            await update_state(f"*Task failed: {exc}*", is_final=True)

    await nc.subscribe("hermes.worker.queue", queue="hermes_workers", cb=message_handler)

    while True:
        await asyncio.sleep(1)


if __name__ == "__main__":
    asyncio.run(run_worker())
